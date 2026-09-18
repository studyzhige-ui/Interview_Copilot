"""Export contract tests. Browser interaction checks are documented in QA.md."""
import sys, unittest, json
from pathlib import Path
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import server
import pymupdf
from docx import Document
from PIL import Image

class Exports(unittest.TestCase):
    def setUp(self):
        self.payload={'document':{'schema':1,'id':'00000000-0000-4000-8000-000000000001','name':'导出验收','sections':[], 'style':{'font':'sans','size':11,'leading':1.5,'margin':18,'gap':17,'template':'modern','color':'#355d82'}},'pages':[]}
        for text in ['中文简历验收','第二页 English 2026']:
            self.payload['pages'].append({'texts':[{'text':text,'font':'sans','bold':True,'size':16,'x':70,'y':80,'height':22,'end':330,'color':[36,48,59]}],'lines':[],'images':[],'columns':[[{'className':'atom paragraph','html':f'<p><strong>{text}</strong> <span style="color: rgb(51, 102, 238)">蓝色文字</span></p>'}]]})
    def test_pdf_pages_and_selectable_chinese(self):
        pdf=pymupdf.open(stream=server.pdf_bytes(self.payload),filetype='pdf')
        self.assertEqual(len(pdf),2)
        self.assertIn('中文简历验收',pdf[0].get_text())
        self.assertIn('English 2026',pdf[1].get_text())
        self.assertAlmostEqual(pdf[0].rect.width,595.5,1)
    def test_docx_real_runs_and_color(self):
        d=Document(BytesIO(server.docx_bytes(self.payload)))
        self.assertIn('中文简历验收',''.join(p.text for p in d.paragraphs))
        self.assertTrue(any(r.bold for p in d.paragraphs for r in p.runs))
        self.assertTrue(any(str(r.font.color.rgb)=='3366EE' for p in d.paragraphs for r in p.runs))
        self.assertIn('w:type="page"',d._element.xml)
    def test_png_contains_both_pages(self):
        im=Image.open(BytesIO(server.png_bytes(self.payload)))
        self.assertEqual(im.width,894)
        self.assertGreater(im.height,2500)
    def test_two_columns_docx_remains_text(self):
        self.payload['pages'][0]['columns'].append([{'className':'atom paragraph','html':'<p>右栏内容</p>'}])
        d=Document(BytesIO(server.docx_bytes(self.payload)))
        self.assertEqual(len(d.tables),1)
        self.assertIn('右栏内容',d.tables[0].cell(0,1).text)
    def test_reject_invalid_document(self):
        self.payload['document']['id']='../../outside'
        with self.assertRaises(ValueError):server.pdf_bytes(self.payload)

if __name__=='__main__': unittest.main()
