"""Local-only functional resume editor. Run with the bundled Python runtime."""
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlsplit, unquote, quote
from io import BytesIO
import json, re, os, base64, threading, traceback, zipfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
PROFILE = ROOT.parent / 'personal-profile'
DATA = ROOT / 'local-data'
LOCK = threading.Lock()
UUID = r'[a-f0-9-]{36}'

def setup_fonts():
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as RLFont
    names = {'sans':'msyh.ttc', 'sans-bold':'msyhbd.ttc', 'serif':'simsun.ttc'}
    metrics = {}
    for key, name in names.items():
        dst = ROOT/'vendor'/f'{key}.ttf'
        if not dst.exists():
            f = TTFont(str(Path('C:/Windows/Fonts')/name),fontNumber=0)
            f.save(str(dst))
        f = TTFont(str(dst))
        h = f['hhea']
        metrics[key] = h.ascent/(h.ascent-h.descent)
        pdfmetrics.registerFont(RLFont(key,str(dst)))
    return metrics

FONT_METRICS = setup_fonts()

def validate_document(d):
    if not isinstance(d,dict) or d.get('schema') != 1 or not re.fullmatch(UUID,str(d.get('id',''))):
        raise ValueError('不支持的简历文档')
    if not isinstance(d.get('sections'),list) or len(d['sections'])>100:
        raise ValueError('简历最多支持 100 个模块')
    return d

def validate_export(p):
    validate_document(p['document'])
    pages=p.get('pages',[])
    if not pages or len(pages)>60: raise ValueError('导出支持 1–60 页')
    if sum(len(p.get('texts',[])) for p in pages)>100000: raise ValueError('文字数量过多')

def pdf_bytes(payload):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    validate_export(payload)
    stream=BytesIO(); canvas=Canvas(stream,pagesize=(595.5,842.25),pageCompression=1)
    canvas.setTitle(payload['document'].get('name','简历'))
    for page in payload['pages']:
        for line in page.get('lines',[]):
            canvas.setFillColorRGB(*[n/255 for n in line['color']])
            canvas.rect(line['x']*.75,(1123-line['y']-line['height'])*.75,line['width']*.75,line['height']*.75,stroke=0,fill=1)
        for run in page['texts']:
            font=run['font'] if run['font'] in FONT_METRICS else 'sans'
            if font=='sans' and run.get('bold'):font='sans-bold'
            size=max(4,min(100,float(run['size'])))*.75
            baseline=(1123-run['y']-run['height']*FONT_METRICS[font])*.75
            canvas.setFillColorRGB(*[n/255 for n in run['color']]);canvas.setFont(font,size)
            canvas.saveState()
            if run.get('italic'):
                canvas.translate(run['x']*.75,baseline);canvas.transform(1,0,.18,1,0,0);canvas.drawString(0,0,run['text'])
            else: canvas.drawString(run['x']*.75,baseline,run['text'])
            canvas.restoreState()
            if run.get('underline') or run.get('strike'):
                canvas.setStrokeColorRGB(*[n/255 for n in run['color']]);canvas.setLineWidth(.5)
                for offset in ([-1.5] if run.get('underline') else [])+([size*.3] if run.get('strike') else []):
                    canvas.line(run['x']*.75,baseline+offset,run['end']*.75,baseline+offset)
        for img in page.get('images',[]):
            if not re.match(r'^data:image/(png|jpeg|webp);base64,',img['src']):continue
            data=base64.b64decode(img['src'].split(',',1)[1])
            canvas.drawImage(ImageReader(BytesIO(data)),img['x']*.75,(1123-img['y']-img['height'])*.75,img['width']*.75,img['height']*.75,mask='auto')
        for link in page.get('links',[]):
            if re.match(r'^(https?://|mailto:)',link['href']):
                canvas.linkURL(link['href'],(link['x']*.75,(1123-link['y']-link['height'])*.75,(link['x']+link['width'])*.75,(1123-link['y'])*.75),relative=0)
        canvas.showPage()
    canvas.save();return stream.getvalue()

def docx_bytes(payload):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from bs4 import BeautifulSoup, NavigableString
    validate_export(payload)
    document=Document();st=payload['document']['style'];font='宋体' if st['font']=='serif' else '微软雅黑'
    section=document.sections[0];section.page_width=Mm(210.08);section.page_height=Mm(297.13)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(st['margin'])
    normal=document.styles['Normal'];normal.font.name=font;normal.font.size=Pt(st['size']);normal._element.rPr.rFonts.set(qn('w:eastAsia'),font)
    normal.paragraph_format.line_spacing=st['leading'];normal.paragraph_format.space_after=Pt(3)
    def put_run(p,node,attrs=None):
        attrs=dict(attrs or {})
        if isinstance(node,NavigableString):
            r=p.add_run(str(node));r.bold=attrs.get('bold');r.italic=attrs.get('italic');r.underline=attrs.get('underline');r.font.strike=attrs.get('strike')
            if attrs.get('color'):
                try:r.font.color.rgb=RGBColor.from_string(attrs['color'])
                except ValueError:pass
            if attrs.get('size'):r.font.size=Pt(attrs['size'])
            r.font.name=font;r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font);return
        if node.name=='br':p.add_run().add_break();return
        if node.name in ['b','strong']:attrs['bold']=True
        if node.name in ['i','em']:attrs['italic']=True
        if node.name=='u':attrs['underline']=True
        if node.name in ['s','strike']:attrs['strike']=True
        style=node.get('style','');color=re.search(r'(?:^|;)\s*color:\s*([^;]+)',style)
        if color:
            v=color.group(1).strip();m=re.match(r'rgb\((\d+),\s*(\d+),\s*(\d+)\)',v)
            if m:attrs['color']=''.join(f'{int(x):02x}' for x in m.groups())
            elif re.fullmatch('#[0-9a-fA-F]{6}',v):attrs['color']=v[1:]
        size=re.search(r'font-size:\s*([\d.]+)(px|em)',style)
        if size:attrs['size']=float(size.group(1))*(.75 if size.group(2)=='px' else st['size'])
        for c in node.children:put_run(p,c,attrs)
    def add_atom(parent,atom):
        soup=BeautifulSoup(atom['html'],'html.parser');cls=atom['className']
        if 'identity' in cls:
            image=soup.find('img')
            if image and str(image.get('src','')).startswith('data:image/'):
                p=parent.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run().add_picture(BytesIO(base64.b64decode(image['src'].split(',',1)[1])),width=Mm(15),height=Mm(20))
            for selector in ['h1','.role','.contact']:
                node=soup.select_one(selector)
                if node:
                    p=parent.add_paragraph();put_run(p,node);p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.font.size=Pt(22.5 if selector=='h1' else 11.25 if selector=='.role' else 8.25);r.bold=selector=='h1'
            return
        p=parent.add_paragraph()
        if 'sectionheading' in cls:
            put_run(p,soup,{'bold':True,'color':st['color'][1:] if st['template']!='classic' else '24303b','size':12})
            p.paragraph_format.space_before=Pt(st['gap']*.75);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True
            if st['template']=='classic':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pr=p._p.get_or_add_pPr();borders=OxmlElement('w:pBdr');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'DBE2EB');borders.append(b);pr.append(borders)
        elif 'entryhead' in cls:
            left=soup.find('strong');right=soup.find('small')
            if left:put_run(p,left,{'bold':True,'size':10.5})
            if right:p.add_run('    ');put_run(p,right,{'size':8.25,'color':'6C7885'})
            p.paragraph_format.keep_with_next=True
        elif 'entrymeta' in cls:
            put_run(p,soup,{'size':9,'color':'627486'});p.paragraph_format.keep_with_next=True
        else:
            put_run(p,soup)
            aligned=soup.find(style=re.compile('text-align'))
            if aligned:
                alignment=re.search(r'text-align:\s*(\w+)',aligned['style'])
                if alignment:p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}.get(alignment.group(1),WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.widow_control=True
    for n,page in enumerate(payload['pages']):
        if n:document.add_page_break()
        columns=page['columns']
        if len(columns)==1:
            for atom in columns[0]:add_atom(document,atom)
        else:
            table=document.add_table(rows=1,cols=2);table.autofit=False
            width=(210-2*st['margin'])/2
            for j,col in enumerate(columns):
                cell=table.cell(0,j);cell.width=Mm(width)
                for atom in col:add_atom(cell,atom)
                if len(cell.paragraphs)>1 and not cell.paragraphs[0].text:cell._element.remove(cell.paragraphs[0]._element)
    document.core_properties.title=payload['document']['name'];document.core_properties.author=''
    stream=BytesIO();document.save(stream);return stream.getvalue()

def png_bytes(payload):
    import pymupdf
    from PIL import Image
    pdf=pymupdf.open(stream=pdf_bytes(payload),filetype='pdf')
    if len(pdf)>12:raise ValueError('长图最多支持 12 页，请改用 PDF')
    ims=[]
    for page in pdf:
        pix=page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5));ims.append(Image.frombytes('RGB',[pix.width,pix.height],pix.samples))
    image=Image.new('RGB',(max(i.width for i in ims),sum(i.height for i in ims)+16*(len(ims)-1)),(237,241,246));y=0
    for im in ims:image.paste(im,(0,y));y+=im.height+16
    stream=BytesIO();image.save(stream,'PNG');return stream.getvalue()

class Handler(SimpleHTTPRequestHandler):
    def log_message(self, fmt, *args): pass
    def send(self,data,ctype='application/json; charset=utf-8',status=200):
        if not isinstance(data,bytes):data=json.dumps(data,ensure_ascii=False).encode('utf-8')
        self.send_response(status);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','no-store');self.end_headers();self.wfile.write(data)
    def safe(self):
        host=self.headers.get('Host','')
        return host in ['127.0.0.1:8790','localhost:8790'] and self.headers.get('Origin','http://'+host)=='http://'+host
    def read_json(self):
        if not self.safe():raise ValueError('仅允许本机同源请求')
        n=int(self.headers.get('Content-Length','0'))
        if not 0<n<=12*1024*1024:raise ValueError('请求大小不受支持')
        return json.loads(self.rfile.read(n))
    def do_GET(self):
        path=urlsplit(self.path).path
        if path=='/api/resumes':
            with LOCK:
                docs=[json.loads(p.read_text(encoding='utf-8')) for p in DATA.glob('*.json') if not p.name.endswith('.history.json')]
            self.send(sorted([{k:d.get(k) for k in ['id','name','createdAt','updatedAt']} for d in docs],key=lambda d:d.get('updatedAt') or '',reverse=True));return
        exported=re.fullmatch('/api/resumes/('+UUID+r')/export\.(pdf|docx|png)',path)
        if exported:
            p=DATA/'exports'/(exported[1]+'.'+exported[2])
            if not p.exists():self.send({'error':'文件不存在'},status=404);return
            data=p.read_bytes();self.send_response(200)
            self.send_header('Content-Type',{'pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','png':'image/png'}[exported[2]])
            record=DATA/(exported[1]+'.json')
            name=json.loads(record.read_text(encoding='utf-8')).get('name','resume') if record.exists() else 'resume'
            self.send_header('Content-Disposition',"attachment; filename=\"resume."+exported[2]+"\"; filename*=UTF-8''"+quote(name+'.'+exported[2]))
            self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data);return
        m=re.fullmatch('/api/resumes/('+UUID+')(/history)?',path)
        if m:
            p=DATA/(m[1]+('.history' if m[2] else '')+'.json')
            if not p.exists():self.send([] if m[2] else {'error':'版本不存在'},status=200 if m[2] else 404);return
            with LOCK:data=p.read_bytes()
            self.send(data);return
        return super().do_GET()
    def translate_path(self,path):
        path=unquote(urlsplit(path).path)
        if path in ['/', '/index.html']:
            base=ROOT.parent/'shared';relative='index.html'
        elif path in ['/today', '/today/']:
            base=ROOT.parent/'today';relative='today-orbit.html'
        elif path.startswith('/today/'):
            base=ROOT.parent/'today';relative=path[len('/today/'):]
        elif path.startswith('/shared/'):
            base=ROOT.parent/'shared';relative=path[len('/shared/'):]
        elif path in ['/profile', '/profile/']:
            base=PROFILE;relative='index.html'
        elif path.startswith('/profile/'):
            base=PROFILE;relative=path[len('/profile/'):]
        elif path.startswith('/market/'):
            base=ROOT.parent/'job-market';relative=path[len('/market/'):] or 'index.html'
        elif path.startswith('/editor/'):
            base=ROOT;relative=path[len('/editor/'):] or 'index.html'
            if relative.startswith(('local-data','tests','__pycache__')):return str(ROOT/'404')
        elif path in ['/editor','/editor/']:
            base=ROOT;relative='index.html'
        else:base=PROFILE;relative=path.lstrip('/') or 'index.html'
        target=(base/relative).resolve()
        if not target.is_relative_to(base.resolve()) or target.suffix not in ['.html','.css','.js','.ttf','.png','.jpg','.svg']:return str(ROOT/'404')
        return str(target)
    def do_PUT(self):
        try:
            m=re.fullmatch('/api/resumes/('+UUID+')',urlsplit(self.path).path)
            if not m:raise ValueError('无效路径')
            d=validate_document(self.read_json())
            if d['id']!=m[1]:raise ValueError('版本 ID 不一致')
            DATA.mkdir(exist_ok=True)
            with LOCK:
                p=DATA/(m[1]+'.json');hp=DATA/(m[1]+'.history.json')
                if p.exists():
                    prev=json.loads(p.read_text(encoding='utf-8'))
                    if prev!=d:
                        hist=json.loads(hp.read_text(encoding='utf-8')) if hp.exists() else []
                        hist.insert(0,prev);hp.write_text(json.dumps(hist[:30],ensure_ascii=False),encoding='utf-8')
                temp=p.with_suffix('.tmp');temp.write_text(json.dumps(d,ensure_ascii=False),encoding='utf-8');os.replace(temp,p)
            self.send({'saved':True})
        except Exception as e:self.send({'error':str(e)},status=400)
    def do_POST(self):
        try:
            if urlsplit(self.path).path=='/api/jobs/parse':
                from job_parser import parse_job
                payload=self.read_json()
                self.send(parse_job(str(payload.get('url',''))));return
            kind=urlsplit(self.path).path.removeprefix('/api/export/')
            if kind not in ['pdf','docx','png']:raise ValueError('不支持的导出格式')
            payload=self.read_json();validate_export(payload)
            func={'pdf':pdf_bytes,'docx':docx_bytes,'png':png_bytes}[kind]
            with LOCK:
                data=func(payload);out=DATA/'exports';out.mkdir(exist_ok=True)
                (out/(payload['document']['id']+'.'+kind)).write_bytes(data)
            self.send({'url':'/api/resumes/'+payload['document']['id']+'/export.'+kind,'bytes':len(data)})
        except Exception as e:
            traceback.print_exc();self.send({'error':str(e)},status=400)

if __name__=='__main__':
    DATA.mkdir(exist_ok=True)
    print('Career OS profile + resume editor: http://127.0.0.1:8790/',flush=True)
    ThreadingHTTPServer(('127.0.0.1',8790),Handler).serve_forever()
